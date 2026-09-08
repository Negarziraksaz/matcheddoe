"""Recreate compact examples from versioned public data. See DATA_SOURCES.json."""

import hashlib
import json
from pathlib import Path

import openpyxl
import pandas as pd

DATA = DEST = Path(".")


def designs():
    workbook = openpyxl.load_workbook(DATA / "v28r4bz6s5" / "LEVO_RSM.xlsx", data_only=True, read_only=True)
    rows = [
        r
        for r in workbook["BBD Model"].values
        if isinstance(r[0], (float, int))
        and r[0] in range(1, 16)
        and all(isinstance(r[i], (float, int)) for i in range(1, 6))
    ]
    pd.DataFrame(
        [r[:5] for r in rows], columns=["sample_id", "dose_g_L", "concentration_ppm", "pH", "response"]
    ).to_csv(DEST / "levo_design.csv", index=False)
    workbook.close()
    workbook = openpyxl.load_workbook(DATA / "87tfjn5s2h" / "Data accessibility.xlsx", data_only=True, read_only=True)
    sheet = workbook["Data and Design"]
    rows = [[sheet.cell(i, j).value for j in range(2, 9)] for i in range(5, 14)]
    pd.DataFrame(
        [[r[0], 36 + 12 * r[1], 50 + 5 * r[2], r[5], r[6]] for r in rows],
        columns=["sample_id", "time_h", "temperature_C", "moisture_percent", "colour_deltaE"],
    ).to_csv(DEST / "banana_design.csv", index=False)
    workbook.close()


def main():
    import argparse
    import urllib.request

    global DATA, DEST
    root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(
        description="Recreate the bundled examples from checksum-verified public measurements."
    )
    parser.add_argument("--raw-dir", type=Path, default=root / "user_data" / "public_raw")
    parser.add_argument("--output", type=Path, default=root / "runs" / "recreated_examples")
    parser.add_argument("--download", action="store_true", help="Fetch missing CC BY 4.0 source files over HTTPS.")
    args = parser.parse_args()
    DATA, DEST = args.raw_dir.resolve(), args.output.resolve()
    if DEST.exists() and any(DEST.iterdir()):
        parser.error("Choose a new or empty output directory; existing files are preserved.")
    sources = json.loads((root / "DATA_SOURCES.json").read_text(encoding="utf-8"))
    for source in sources:
        if source["license"] != "CC-BY-4.0":
            parser.error("This adapter only accepts the reviewed CC BY 4.0 source manifest.")
        for item in source["downloads"]:
            path = DATA / source["repository_id"] / item["file"]
            if not path.exists():
                if not args.download:
                    parser.error(f"Missing {path}. Supply --download or place the original file there.")
                print(f"Downloading {source['doi']} ({source['license']}): {item['file']}")
                request = urllib.request.Request(
                    item["url"], headers={"User-Agent": "Research-example-reproduction/0.1"}
                )
                with urllib.request.urlopen(request, timeout=60) as response:
                    data = response.read()
                if len(data) != item["bytes"] or hashlib.sha256(data).hexdigest() != item["sha256"]:
                    parser.error(f"Download checksum mismatch for {item['file']}; no source file was saved.")
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_bytes(data)
            data = path.read_bytes()
            if len(data) != item["bytes"] or hashlib.sha256(data).hexdigest() != item["sha256"]:
                parser.error(
                    f"Source checksum mismatch for {path}. Check the dataset version; the file was not modified."
                )
    DEST.mkdir(parents=True, exist_ok=True)
    designs()
    manifest = {
        p.name: {"sha256": hashlib.sha256(p.read_bytes()).hexdigest(), "bytes": p.stat().st_size}
        for p in DEST.iterdir()
        if p.is_file() and p.name != "manifest.json"
    }
    (DEST / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    print(
        f"Examples recreated in {DEST}. Compare numeric CSV values with examples/; float formatting can vary with library versions."
    )


if __name__ == "__main__":
    main()
